import survey_ocr as s, grid_pipeline as gp, json, importlib
from openpyxl import load_workbook
import make_dataentry_csv as mk
importlib.reload(gp)
cfg=json.load(open('form_config.json')); fields=cfg['fields']; gcfg=cfg.get('groups',{})
cal,meta=gp.calibrate_blank('BRNB422005F2960_004268.pdf',fields,scan_rotation=90,overrides=gp.load_overrides('calibration_overrides.json'))
CK=load_workbook('/sessions/wonderful-elegant-newton/mnt/uploads/Multiple_Choice_checked.xlsx')
qmap={q:kind for q,(txt,kind) in mk.QUESTIONS.items()}
# reverse option map: full text -> short label (per group)
rev={grp:{full:short for short,full in d.items()} for grp,d in mk.OPTIONS.items()}
def norm_truth(grp,val):
    if val is None: return None
    v=str(val).split('-')[0].strip()  # strip specify suffix
    # match against full option texts
    for short,full in mk.OPTIONS.get(grp,{}).items():
        if full.lower().startswith(v.lower()) or v.lower().startswith(full.split('(')[0].strip().lower()):
            return short
    return v  # fallback
FILES={'31':'diag_it31.pdf','32':'diag_it32.pdf','33':'diag_it33.pdf'}
# page offsets per serial (from processing order)
def survey_pages(pdf, n):
    return [{p:s.pdf_page_to_image(pdf,1+i*2+p,dpi=300,scan_rotation=90,color=True) for p in range(2)} for i in range(n)]
tot_wrong=tot_miss=fixed_wrong=fixed_miss=0
details=[]
for sh,pdf in FILES.items():
    ws=CK[sh]; serials=[str(ws.cell(2,c).value) for c in range(3,ws.max_column+1)]
    pagesets=survey_pages(pdf,len(serials))
    codeout=[]
    for pg in pagesets:
        w,_=gp.read_survey(pg,fields,cal,meta,gcfg); codeout.append(w)
    for r in range(3,ws.max_row+1):
        qn=ws.cell(r,1).value; grp=qmap.get(qn)
        if not grp or grp=='text': continue
        for ci,c in enumerate(range(3,ws.max_column+1)):
            cell=ws.cell(r,c)
            fg=cell.fill.fgColor.rgb if cell.fill and cell.fill.patternType else None
            fontc=cell.font.color.rgb if cell.font and cell.font.color else None
            yellow= fg=='FFFFFF00'; red= fontc in ('FFFF0000','FF0000')
            if not (yellow or red): continue
            truth=norm_truth(grp,cell.value)
            code=codeout[ci].get(grp)
            got = (code==truth)
            if yellow: tot_miss+=1; fixed_miss+=got
            if red: tot_wrong+=1; fixed_wrong+=got
            if not got: details.append((sh,serials[ci],qn,truth,code,'MISS' if yellow else 'WRONG'))
print(f'WRONG cells: {fixed_wrong}/{tot_wrong} now correct')
print(f'MISSING cells: {fixed_miss}/{tot_miss} now correct')
print(f'TOTAL flagged fixed: {fixed_wrong+fixed_miss}/{tot_wrong+tot_miss}')
